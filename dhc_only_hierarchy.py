#!/usr/bin/env python3
"""Definitive-only hierarchy over the Definitive ids Zeus holds.

The base is every Zeus entity that carries a Definitive id, with the id
exactly as Zeus stores it (DHC_Id), whatever its verdict:
  - the accuracy run's _scored and _unverifiable files together (the audited
    entities), and
  - every active Definitive-import-created entity with a Definitive id ("Zeus
    import entities.sql", `zeus.import_query_file`), which the audits exclude.
    Linked or not: the query reads no link table. Its id is not audited, so
    its Verdict is "Import-created (not audited)".
Base_Source says which. Everything above the base comes from
Definitive's own ownership data, the hierarchy views dhc_hierarchy.py walks,
with the same parent rule (SfParentAccountId, else IdNetwork, else the record
itself). Nothing comes from Zeus's own hierarchy: no LinkClientWorkLocation,
no LinkHealthSystemClient, no Zeus role deciding a parent, and no audit
correction or coverage proposal replacing the id Zeus holds. Compare
zeus_hierarchy.py, which blends the two.

The tree holds every Definitive record a Zeus entity points at plus every
owner above it, up to the top-level owner. A parent Definitive names that is
in no hierarchy view (a corporate owner such as The US Oncology Network) gets
a row of its own, so every Immediate_ParentId in the tree has a row.

  py dhc_only_hierarchy.py --config sources.yaml
      --accuracy "Results Output/<accuracy run>/<run>_scored.csv"
      [--definitive-from "Results Output/<dhc_hierarchy run>"]
      [--zeus-imports "Results Output/<earlier dhc_only_hierarchy run>"]
      [--no-imports] [--label <suffix>]

Without --zeus-imports the import query runs live against Zeus (ReadOnly);
either way it is snapshotted as <run>_zeus_imports.csv. --no-imports leaves
the import-created entities out (the first, audited-only form).

An id that cannot be placed is classed from the accuracy run's own
Definitive snapshots (_dhc_<snapshot>.parquet): a GPO, a practice location
parent, a record in an overview but not in the hierarchy views, or in no
Definitive source.

Writes into "Results Output/dhc_only_hierarchy_<YYYY_MM_DD_HHMM>/":
  <run>_entities.csv            one row per Zeus entity: its id, the record's
                                immediate and ultimate parent, and the Zeus
                                entities that hold those parents' ids
  <run>_hierarchy.csv           one row per Definitive record in the tree
  <run>_hierarchy_immediate.csv edge list: record -> direct owner
  <run>_hierarchy_ultimate.csv  edge list: record -> top-level owner
  <run>_zeus_imports.csv        the import-created entities, as queried
  <run>_dhc_*hierarchy.parquet/.sql
  Definitive_Only_Hierarchy_<run date and time>.xlsx

Without --definitive-from the hierarchy views are queried live.
"""
import argparse
import os
import shutil
import warnings

import pandas as pd
from openpyxl import Workbook

from dhc_match_v2 import RESULTS_DIR, load_config, new_run_prefix, \
    zeus_connect
from dhc_hierarchy import LABELS, SEP, _xl, labelled, load_hierarchy, \
    print_checks, resolve_hierarchy_prefix, walk
from build_audit_workbook import FILL_NAVY, FILL_PURPLE, H1, H2, SECTION, \
    WRAP, _put, _table, build_methodology, sheet_data, workbook_path

UNVERIFIABLE = 'Unverifiable - id in no Definitive source'
IMPORTED = 'Import-created (not audited)'
EXTERNAL_RULE = 'Not in any hierarchy view'
IMPORT_SNAPSHOT = 'zeus_imports'

# Id_Status: where the id Zeus holds sits in the Definitive hierarchy.
IN_VIEW = 'Record in a hierarchy view'
OWNER_ONLY = 'Owner named by Definitive, in no hierarchy view'
GPO = 'GPO - not in the ownership hierarchy'
LOCATION_ONLY = 'Practice location parent only - in no hierarchy view'
OVERVIEW_ONLY = 'In a Definitive overview, not in the hierarchy views'
NO_SOURCE = 'In no Definitive source'

ENTITY_COLUMNS = [
    'EntityId', 'Base_Source', 'Import_Description', 'Zeus_Name',
    'Zeus_Sources', 'Zeus_City', 'Zeus_State',
    'DHC_Id', 'DHC_Id_Source', 'DHC_Id_Conflict', 'Verdict', 'Geo_Conflict',
    'Id_Status', 'Definitive_Name', 'TypeFirm', 'TypeHospital',
    'CompanyStatus', 'Source_View', 'Is_Root',
    'Immediate_ParentId', 'Immediate_ParentName', 'Immediate_Parent_TypeFirm',
    'Immediate_Parent_Zeus_EntityIds', 'Immediate_Level',
    'Ultimate_ParentId', 'Ultimate_ParentName', 'Ultimate_Parent_TypeFirm',
    'Ultimate_Parent_Zeus_EntityIds', 'Has_Intermediate_Parent',
    'Path_Ids', 'Path_Names', 'Hierarchy_Issue']
TREE_COLUMNS = [
    'HospitalId', 'HospitalName', 'TypeFirm', 'TypeHospital', 'CompanyStatus',
    'HqCity', 'HqState', 'Source_View', 'Node_Role', 'Is_Root',
    'ParentId', 'ParentName', 'Parent_TypeFirm', 'Parent_Rule',
    'Parent_In_Definitive', 'Level', 'Child_Count',
    'RootId', 'RootName', 'Root_TypeFirm', 'Ultimate_Level',
    'Has_Intermediate_Parent', 'Descendant_Count',
    'Definitive_Descendant_Count', 'Zeus_Entity_Count', 'Zeus_EntityIds',
    'Zeus_Sources', 'Zeus_Verdicts', 'Subtree_Zeus_Entity_Count',
    'Tree_Zeus_Entity_Count', 'Path_Ids', 'Path_Names', 'Hierarchy_Issue']
EDGE_IMMEDIATE = ['HospitalId', 'HospitalName', 'TypeFirm', 'ParentId',
                  'ParentName', 'Parent_TypeFirm', 'Level', 'Is_Root',
                  'Parent_Rule', 'Node_Role', 'Zeus_EntityIds']
EDGE_ULTIMATE = ['HospitalId', 'HospitalName', 'TypeFirm', 'RootId',
                 'RootName', 'Root_TypeFirm', 'Ultimate_Level', 'Is_Root',
                 'Has_Intermediate_Parent', 'Node_Role', 'Zeus_EntityIds']


def _ids(s):
    return pd.to_numeric(s, errors='coerce').astype('Int64')


def _join(v):
    return '|'.join(sorted({str(x) for x in v if pd.notna(x)}, key=int))


# ============================================================================
# The base: every Zeus entity with a Definitive id
# ============================================================================
def load_base(scored):
    """One row per Zeus entity carrying a Definitive id, from both halves of
    an accuracy run. Returns (frame, rows in _scored, rows in _unverifiable)."""
    tail = '_scored.csv'
    if not scored.endswith(tail):
        raise SystemExit(f'{scored} should end in {tail}')
    unv = scored[:-len(tail)] + '_unverifiable.csv'
    if not os.path.exists(unv):
        raise SystemExit(f'Missing {unv}: it is the other half of the '
                         f'accuracy run, so keep the run folder intact.')
    s = pd.read_csv(scored, low_memory=False).assign(Base_Source='Accuracy')
    u = pd.read_csv(unv, low_memory=False).assign(
        Base_Source='Accuracy (unverifiable)', Verdict=UNVERIFIABLE,
        DHC_Id_Conflict=pd.NA, Geo_Conflict=pd.NA)
    z = pd.concat([s[BASE], u[BASE]], ignore_index=True)
    return z, len(s), len(u)


BASE = ['EntityId', 'Base_Source', 'Zeus_Name', 'Zeus_Sources', 'Zeus_City',
        'Zeus_State', 'DHC_Id', 'DHC_Id_Source', 'DHC_Id_Conflict', 'Verdict',
        'Geo_Conflict']
IMPORT_ROLES = ['Client', 'HealthSystem', 'WorkLocation']   # Zeus_Sources order


def load_imports(zc, prefix, replay=None):
    """Every active import-created entity with a Definitive id, live or
    replayed, always written to this run's folder and read back from it, so
    a live run and its replay see byte-identical input."""
    out = f'{prefix}_{IMPORT_SNAPSHOT}.csv'
    if replay:
        src = replay
        if os.path.isdir(src):
            hits = [f for f in os.listdir(src)
                    if f.endswith(f'_{IMPORT_SNAPSHOT}.csv')]
            if len(hits) != 1:
                raise SystemExit(f'{src} holds {len(hits)} *_{IMPORT_SNAPSHOT}'
                                 f'.csv files; pass the file instead.')
            src = os.path.join(src, hits[0])
        print(f'Zeus imports: replayed from {src}')
        shutil.copyfile(src, out)          # a copy, so the snapshot is exact
    else:
        if not zc.get('import_query_file'):
            raise SystemExit('zeus config needs import_query_file, or pass '
                             '--no-imports.')
        with zeus_connect(zc.get('connection') or {}) as cx, \
                warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)   # pandas vs pyodbc
            df = pd.read_sql(open(zc['import_query_file']).read(), cx)
        df.to_csv(out, index=False)
        print(f'  {IMPORT_SNAPSHOT:20} {len(df):>7,} rows  '
              f'({os.path.basename(zc["import_query_file"])})')
    imp = pd.read_csv(out, low_memory=False)
    ent = _ids(imp.Entity_DHC_VerifiedSourceId)
    levs = _ids(imp.LEVS_DHC_VerifiedSourceId)
    flags = {r: imp[f'Is{r}'].astype(str).str.strip().str.lower()
             .isin(['true', '1', '1.0']) for r in IMPORT_ROLES}
    return pd.DataFrame({
        'EntityId': imp.EntityId.astype(int),
        'Base_Source': 'Import-created',
        'Import_Description': imp.EntityDescription,
        'Zeus_Name': imp.Entity_Name.fillna(imp.ClientInfoName)
                     .fillna(imp.WorkLocationInfoName)
                     .fillna(imp.HealthSystemInfoName),
        'Zeus_Sources': ['|'.join(r for r in IMPORT_ROLES if flags[r].iloc[i])
                         for i in range(len(imp))],
        'Zeus_City': imp.City, 'Zeus_State': imp.State,
        # As the audits choose: Entity's own id first, else LinkEntityVerified
        # Source; a conflict is either column disagreeing with the other.
        'DHC_Id': ent.fillna(levs),
        'DHC_Id_Source': ent.notna().map({True: 'Entity_DHC_VerifiedSourceId',
                                          False: 'LEVS_DHC_VerifiedSourceId'}),
        'DHC_Id_Conflict': (ent.notna() & levs.notna() & ent.ne(levs))
                           .fillna(False).astype(bool),
        'Verdict': IMPORTED, 'Geo_Conflict': pd.NA})


def check_base(z):
    z['EntityId'] = z.EntityId.astype(int)
    z['DHC_Id'] = _ids(z.DHC_Id)
    if z.EntityId.duplicated().any():
        dup = z[z.EntityId.duplicated(keep=False)]
        raise SystemExit(f'{dup.EntityId.nunique():,} EntityId(s) appear more '
                         f'than once in the base, from '
                         f'{sorted(dup.Base_Source.unique())}.')
    if z.DHC_Id.isna().any():
        raise SystemExit(f'{z.DHC_Id.isna().sum():,} base row(s) carry no '
                         f'numeric DHC_Id; the base must be ids Zeus holds.')
    return z


def definitive_ids(cfg, scored):
    """The ids in each Definitive source, from the accuracy run's own
    snapshots: {'GPO': ..., 'overview': ..., 'location': ...}. Used only to
    say why an id is not in the hierarchy."""
    pre = scored[:-len('_scored.csv')]
    out = {'GPO': set(), 'overview': set(), 'location': set()}
    for kind, blocks in (('definitive', cfg.get('definitive') or []),
                         ('locations', cfg.get('locations') or [])):
        for b in blocks:
            p = f'{pre}_dhc_{b["snapshot"]}.parquet'
            if not os.path.exists(p):
                raise SystemExit(f'Missing {p}: the accuracy run\'s Definitive '
                                 f'snapshot, used to class unplaced ids.')
            ids = set(_ids(pd.read_parquet(p, columns=[b['id']])[b['id']])
                      .dropna().astype(int))
            key = ('location' if kind == 'locations' else
                   'GPO' if b.get('entity_type') == 'GPO' else 'overview')
            out[key] |= ids
    return out


# ============================================================================
# The tree: base records and every owner above them, Definitive only
# ============================================================================
def external_rows(h, ids):
    """A root row for each parent id Definitive names but no view holds."""
    name = (h[~h.Parent_In_Definitive].drop_duplicates('ParentId')
            .set_index('ParentId').ParentName)
    ids = pd.array(sorted(ids), dtype='Int64')
    nm = pd.Series(ids).map(name)
    return pd.DataFrame({
        'HospitalId': ids, 'HospitalName': nm, 'TypeFirm': None,
        'Source_View': '(none)', 'Is_Root': True,
        'ParentId': ids, 'ParentName': nm, 'Parent_TypeFirm': None,
        'Parent_Rule': EXTERNAL_RULE, 'Parent_In_Definitive': False,
        'Level': 0, 'RootId': ids, 'RootName': nm, 'Root_TypeFirm': None,
        'Ultimate_Level': 0, 'Has_Intermediate_Parent': False,
        'Path_Ids': pd.Series(ids).astype(str),
        'Path_Names': nm.fillna('').astype(str),
        'Hierarchy_Issue': 'Not_In_Hierarchy_Views'})


def build_tree(h, z):
    """The Definitive records the base points at, plus their owners.

    Returns (tree, external ids). `h` is the walked hierarchy (dhc_hierarchy
    .walk); each record's Path_Ids already lists its owners root first,
    including an owner in no view, so the tree is the union of the base
    records' paths. Counts are recomputed within this tree;
    Definitive_Descendant_Count keeps the whole-Definitive figure."""
    rec = set(h.HospitalId.dropna().astype(int))
    ext = set(h.loc[~h.Parent_In_Definitive, 'ParentId'].astype(int))
    base = set(z.DHC_Id.astype(int))
    paths = h[h.HospitalId.isin(base & rec)].Path_Ids
    nodes = {int(x) for p in paths for x in p.split(SEP)} | (base & ext)
    t = pd.concat([h[h.HospitalId.isin(nodes & rec)],
                   external_rows(h, nodes & ext)], ignore_index=True)
    t = t.rename(columns={'Descendant_Count': 'Definitive_Descendant_Count'})
    t['Definitive_Descendant_Count'] = t.Definitive_Descendant_Count \
        .fillna(0).astype(int)

    kids = t[~t.Is_Root].groupby('ParentId').size()
    t['Child_Count'] = t.HospitalId.map(kids).fillna(0).astype(int)
    above = pd.Series([int(x) for p in t.Path_Ids
                       for x in p.split(SEP)[:-1]]).value_counts()
    t['Descendant_Count'] = t.HospitalId.map(above).fillna(0).astype(int)

    g = z.groupby('DHC_Id').agg(
        Zeus_Entity_Count=('EntityId', 'size'),
        Zeus_EntityIds=('EntityId', _join),
        Zeus_Sources=('Zeus_Sources', lambda v: '|'.join(sorted(
            {x for s in v.dropna() for x in str(s).split('|')}))),
        Zeus_Verdicts=('Verdict', lambda v: '|'.join(sorted(set(v.dropna())))))
    t = t.merge(g, left_on='HospitalId', right_index=True, how='left')
    t['Zeus_Entity_Count'] = t.Zeus_Entity_Count.fillna(0).astype(int)
    t['Node_Role'] = t.Zeus_Entity_Count.gt(0).map(
        {True: 'Zeus id', False: 'Owner only'})
    # Zeus entities at or below each node: each entity counted once on every
    # node of its id's path.
    sub = pd.Series(
        [n for p, c in zip(t.Path_Ids, t.Zeus_Entity_Count) if c
         for n in [int(x) for x in p.split(SEP)] for _ in range(c)]
    ).value_counts()
    t['Subtree_Zeus_Entity_Count'] = t.HospitalId.map(sub).fillna(0).astype(int)
    t['Tree_Zeus_Entity_Count'] = t.RootId.map(
        t.groupby('RootId').Zeus_Entity_Count.sum()).fillna(0).astype(int)
    for c in ('TypeHospital', 'CompanyStatus', 'HqCity', 'HqState'):
        if c not in t:
            t[c] = None
    return t[TREE_COLUMNS], nodes & ext


def id_status(ids, T, dids):
    """Why each id is, or is not, in the tree; the first that applies."""
    rule = ids.map(T.Parent_Rule)
    s = pd.Series(NO_SOURCE, index=ids.index)
    for label, pool in ((LOCATION_ONLY, dids['location']),
                        (GPO, dids['GPO']), (OVERVIEW_ONLY, dids['overview'])):
        s[ids.isin(pool)] = label
    s[ids.isin(T.index)] = IN_VIEW
    s[rule.eq(EXTERNAL_RULE)] = OWNER_ONLY
    return s


def place_entities(z, t, dids):
    """Each Zeus entity with its id's position in the tree."""
    T = t.set_index('HospitalId')
    held = T.Zeus_EntityIds
    e = z.copy()
    if 'Import_Description' not in e:
        e['Import_Description'] = None
    e['Id_Status'] = id_status(e.DHC_Id, T, dids)
    for out, col in [('Definitive_Name', 'HospitalName'),
                     ('TypeFirm', 'TypeFirm'), ('TypeHospital', 'TypeHospital'),
                     ('CompanyStatus', 'CompanyStatus'),
                     ('Source_View', 'Source_View'), ('Is_Root', 'Is_Root'),
                     ('Immediate_ParentId', 'ParentId'),
                     ('Immediate_ParentName', 'ParentName'),
                     ('Immediate_Parent_TypeFirm', 'Parent_TypeFirm'),
                     ('Immediate_Level', 'Level'),
                     ('Ultimate_ParentId', 'RootId'),
                     ('Ultimate_ParentName', 'RootName'),
                     ('Ultimate_Parent_TypeFirm', 'Root_TypeFirm'),
                     ('Has_Intermediate_Parent', 'Has_Intermediate_Parent'),
                     ('Path_Ids', 'Path_Ids'), ('Path_Names', 'Path_Names'),
                     ('Hierarchy_Issue', 'Hierarchy_Issue')]:
        e[out] = e.DHC_Id.map(T[col])
    for c in ('Immediate_ParentId', 'Ultimate_ParentId', 'Immediate_Level'):
        e[c] = e[c].astype('Int64')
    # A root is its own parent; listing its own Zeus entities as its parent's
    # would read as a link that is not there.
    owned = e.Is_Root.eq(False)
    e['Immediate_Parent_Zeus_EntityIds'] = e.Immediate_ParentId.map(held) \
        .where(owned)
    e['Ultimate_Parent_Zeus_EntityIds'] = e.Ultimate_ParentId.map(held) \
        .where(owned)
    return e.sort_values(['Ultimate_ParentId', 'Path_Ids', 'EntityId'],
                         na_position='last')[ENTITY_COLUMNS]


# ============================================================================
# Checks and outputs
# ============================================================================
def checks(z, e, t, counts):
    """`counts` is [(base source, rows read)], in the order loaded."""
    ids = set(t.HospitalId.astype(int))
    roots = t[t.Is_Root]
    placed = e[e.Immediate_ParentId.notna()]
    return [
        (f'entities = {" + ".join(k for k, _ in counts)} rows', len(z),
         sum(n for _, n in counts)),
        ('placed + not in the hierarchy = entities',
         len(placed) + int(e.Immediate_ParentId.isna().sum()), len(e)),
        ('sum of Zeus_Entity_Count over the tree = placed entities',
         int(t.Zeus_Entity_Count.sum()), len(placed)),
        ('tree nodes with their parent in the tree = nodes',
         int(t.ParentId.astype(int).isin(ids).sum()), len(t)),
        ('roots + descendants = nodes',
         int(len(roots) + roots.Descendant_Count.sum()), len(t)),
        ('nodes with a Zeus entity at or below = nodes',
         int(t.Subtree_Zeus_Entity_Count.gt(0).sum()), len(t)),
        ('level 0 nodes = roots', int((t.Level == 0).sum()), len(roots)),
        ('nodes in a cycle = 0', int((t.Hierarchy_Issue == 'Cycle').sum()), 0),
    ]


def ultimate_edges(t):
    u = labelled(t, EDGE_ULTIMATE)
    mid = t.sort_values(['RootId', 'Path_Ids'])
    u['Intermediate_ParentId'] = mid.ParentId.where(
        mid.Has_Intermediate_Parent).astype('Int64')
    u['Intermediate_ParentName'] = mid.ParentName.where(
        mid.Has_Intermediate_Parent)
    return u


def build_workbook(z, e, t, rows, path, run, accuracy, imports):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('Summary')
    for col, w in zip('ABCDEF', (4, 52, 16, 16, 16, 60)):
        ws.column_dimensions[col].width = w
    ws.merge_cells('A1:F1')
    _put(ws, 'A1', 'Jackson and Coker Locum Tenens - Definitive-Only '
                   'Hierarchy of Zeus Definitive Ids', H1, FILL_NAVY)
    ws.merge_cells('A2:F2')
    _put(ws, 'A2', f'Run {run} - base: {os.path.basename(accuracy)} and its '
                   f'_unverifiable.csv'
                   + (', plus every import-created entity with an id'
                      if imports else ''), H2, FILL_PURPLE)

    placed = e[e.Immediate_ParentId.notna()]
    roots = t[t.Is_Root]
    multi = roots[roots.Descendant_Count > 0]
    owned = placed[placed.Is_Root.eq(False)]
    n_imp = int(e.Base_Source.eq('Import-created').sum())
    _put(ws, 'B4', 'Headline', SECTION)
    ws.merge_cells('B5:F9')
    _put(ws, 'B5',
         f'{len(e):,} Zeus entities carry a Definitive id'
         + (f' ({len(e) - n_imp:,} audited, {n_imp:,} created by a Definitive '
            f'import)' if imports else '') +
         f'. {len(placed):,} of '
         f'those ids are in the Definitive ownership hierarchy and '
         f'{len(e) - len(placed):,} are not (a GPO, a practice location '
         f'parent, or an id in no Definitive source). The placed entities hold '
         f'{t.Node_Role.eq("Zeus id").sum():,} distinct Definitive records. '
         f'Walking each one up Definitive\'s own ownership chain adds '
         f'{t.Node_Role.eq("Owner only").sum():,} owners that no Zeus entity '
         f'points at, for a tree of {len(t):,} records in {len(roots):,} '
         f'trees, {len(multi):,} of them with more than one member, at most '
         f'{t.Level.max()} levels deep. {len(owned):,} placed entities have a '
         f'Definitive owner above them; the other '
         f'{len(placed) - len(owned):,} are top-level records. No Zeus link, '
         f'Zeus role or audit correction is used anywhere.', align=WRAP)

    _put(ws, 'B11', 'Where each Zeus id sits', SECTION)
    srcs = list(dict.fromkeys(e.Base_Source))
    st = e.groupby(['Id_Status', 'Base_Source']).size().unstack(fill_value=0)
    st = st.loc[e.Id_Status.value_counts().index]
    r = _table(ws, 12, ['Id_Status', 'Entities'] + srcs,
               [[k, int(x.sum())] + [int(x.get(c, 0)) for c in srcs]
                for k, x in st.iterrows()])
    _put(ws, f'B{r}', 'Placed entities by audit verdict', SECTION)
    vt = placed.assign(Owned=placed.Is_Root.eq(False))
    r = _table(ws, r + 1, ['Verdict on the id Zeus holds', 'Entities',
                           'With an owner above', 'Top-level record'],
               [[k, len(g), int(g.Owned.sum()), int((~g.Owned).sum())]
                for k, g in vt.groupby('Verdict')])
    _put(ws, f'B{r}', 'Tree by level', SECTION)
    types = t.TypeFirm.fillna('(not in a view)').value_counts()
    r = _table(ws, r + 1, ['Immediate_Level', 'Records', 'Zeus ids',
                           'Owner only'] + list(types.index),
               [[int(lv), len(g), int(g.Node_Role.eq('Zeus id').sum()),
                 int(g.Node_Role.eq('Owner only').sum())]
                + [int((g.TypeFirm.fillna('(not in a view)') == x).sum())
                   for x in types.index]
                for lv, g in t.groupby('Level')])
    _put(ws, f'B{r}', 'Largest ultimate parents by Zeus entities', SECTION)
    big = multi.sort_values('Tree_Zeus_Entity_Count', ascending=False).head(20)
    r = _table(ws, r + 1, ['Ultimate parent', 'Ultimate_ParentId',
                           'Zeus entities in tree', 'Records in tree',
                           'Records in Definitive tree'],
               [[x.HospitalName, int(x.HospitalId), int(x.Tree_Zeus_Entity_Count),
                 int(x.Descendant_Count) + 1,
                 int(x.Definitive_Descendant_Count) + 1]
                for x in big.itertuples()])
    _put(ws, f'B{r}', 'Identity checks', SECTION)
    _table(ws, r + 1, ['Check', 'Left', 'Right', 'Result'],
           [[k, a, b, 'OK' if a == b else 'FAIL'] for k, a, b in rows])

    build_methodology(wb, [
        ('Base', 'Every Zeus entity that carries a Definitive id, from the '
                 'accuracy run named on the Summary: its _scored.csv and its '
                 '_unverifiable.csv, which together are every in-scope entity '
                 'with an id. The id is the one Zeus stores (DHC_Id), '
                 'whatever the audit verdict on it; Verdict is carried so a '
                 'reader can filter. Recommended corrections and coverage '
                 'proposals are not used.'
                 + (' Also every active entity created by a Definitive import '
                    '(EntityDescription Definitive Physician Group / Provider '
                    '/ Health System Import) with a Definitive-sourced id, '
                    'from "Zeus import entities.sql", linked or not. The '
                    'audits exclude these, so their Verdict is "Import-created '
                    '(not audited)". Base_Source says which part of the base a '
                    'row came from.' if imports else
                    ' Import-created entities are not included (--no-imports).')),
        ('Id_Status', 'Why an id is, or is not, placed: a record in a '
                      'hierarchy view; an owner Definitive names that is in no '
                      'view; or, if unplaced, a record in a Definitive '
                      'overview but not the hierarchy views, a GPO, a '
                      'practice location parent, or in no Definitive source, '
                      'checked in that order against the accuracy run\'s own '
                      'Definitive snapshots.'),
        ('No Zeus hierarchy', 'No Zeus link table (LinkClientWorkLocation, '
                              'LinkHealthSystemClient), Zeus role or Zeus '
                              'parent decides anything here. Every parent is '
                              'Definitive\'s. zeus_hierarchy.py is the blended '
                              'version.'),
        ('Parent rule', 'As in dhc_hierarchy.py: SfParentAccountId if present, '
                        'otherwise IdNetwork, otherwise the record itself, '
                        'which makes it a root. Defined in each view\'s '
                        '"Definitive * Hierarchy.sql"; the SQL that ran is '
                        'saved beside each snapshot.'),
        ('What is in the tree', 'Each Definitive record a Zeus entity points '
                                'at (Node_Role = Zeus id) and every owner '
                                'above it up to the top-level owner (Owner '
                                'only where no Zeus entity points at it). '
                                'Definitive records below the base, and '
                                'practice locations, are not included.'),
        ('Owners in no view', 'A parent Definitive names that is in no '
                              'hierarchy view - a corporate owner such as The '
                              'US Oncology Network - is given its own root '
                              'row, Parent_Rule "Not in any hierarchy view", '
                              'so every parent in the tree has a row.'),
        ('Immediate and ultimate parent', 'Immediate_* is the direct owner, '
                                          'Ultimate_* the top-level owner, as '
                                          'in the Definitive Ownership '
                                          'Hierarchy. A root is its own '
                                          'immediate and ultimate parent.'),
        ('Parent Zeus entities', 'On Entities, Immediate_Parent_Zeus_EntityIds '
                                 'and Ultimate_Parent_Zeus_EntityIds list the '
                                 'Zeus entities whose own Definitive id is '
                                 'that parent. This is the entity-to-entity '
                                 'hierarchy Definitive implies. Blank where '
                                 'the entity\'s record is a root, or where no '
                                 'Zeus entity holds the parent\'s id.'),
        ('Counts', 'Child_Count and Descendant_Count are within this tree. '
                   'Definitive_Descendant_Count is the record\'s whole '
                   'Definitive tree below it. Subtree_Zeus_Entity_Count is '
                   'the Zeus entities at or below a record; '
                   'Tree_Zeus_Entity_Count is the whole tree\'s.'),
        ('Reproducibility', 'Definitive views and Zeus are both live. Pass '
                            '--definitive-from and --zeus-imports with this '
                            'run folder, and the same --accuracy file, to '
                            'regenerate it exactly.'),
    ])
    sheet_data(wb, 'Entities', _xl(e))
    sheet_data(wb, 'Hierarchy', _xl(labelled(t)))
    sheet_data(wb, 'Immediate_Parent', _xl(labelled(t, EDGE_IMMEDIATE)))
    sheet_data(wb, 'Ultimate_Parent', _xl(ultimate_edges(t)))
    sheet_data(wb, 'Roots', _xl(multi.sort_values('Tree_Zeus_Entity_Count',
                                                  ascending=False)
                                .rename(columns=LABELS)))
    sheet_data(wb, 'Not_In_Hierarchy', _xl(e[e.Immediate_ParentId.isna()]))
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--config', required=True)
    ap.add_argument('--accuracy', metavar='SCORED_CSV', required=True,
                    help='a scored accuracy run; it and its _unverifiable.csv '
                         'are the base')
    ap.add_argument('--definitive-from', metavar='PREFIX',
                    help='replay the hierarchy snapshot of an earlier '
                         'dhc_hierarchy or dhc_only_hierarchy run')
    ap.add_argument('--zeus-imports', metavar='RUN_OR_CSV',
                    help='replay the import-created entities of an earlier '
                         'dhc_only_hierarchy run (its folder or _zeus_imports'
                         '.csv) instead of querying Zeus')
    ap.add_argument('--no-imports', action='store_true',
                    help='leave import-created entities out of the base')
    ap.add_argument('--label', help='optional suffix for the run folder name')
    ap.add_argument('--results-dir', default=RESULTS_DIR)
    a = ap.parse_args()
    if a.no_imports and a.zeus_imports:
        raise SystemExit('--no-imports and --zeus-imports contradict.')

    cfg = load_config(a.config)
    z, n_scored, n_unv = load_base(a.accuracy)
    counts = [('scored', n_scored), ('unverifiable', n_unv)]
    dids = definitive_ids(cfg, a.accuracy)
    src = resolve_hierarchy_prefix(cfg, a.definitive_from)
    prefix = new_run_prefix('dhc_only_hierarchy', a.label, a.results_dir)
    if not a.no_imports:
        imp = load_imports(cfg.get('zeus') or {}, prefix, a.zeus_imports)
        z = pd.concat([z, imp], ignore_index=True)
        counts.append(('import-created', len(imp)))
    z = check_base(z)
    h = walk(load_hierarchy(cfg, prefix, src))
    t, ext = build_tree(h, z)
    e = place_entities(z, t, dids)

    e.to_csv(f'{prefix}_entities.csv', index=False)
    labelled(t).to_csv(f'{prefix}_hierarchy.csv', index=False)
    labelled(t, EDGE_IMMEDIATE).to_csv(f'{prefix}_hierarchy_immediate.csv',
                                       index=False)
    ultimate_edges(t).to_csv(f'{prefix}_hierarchy_ultimate.csv', index=False)

    placed = e[e.Immediate_ParentId.notna()]
    roots = t[t.Is_Root]
    print(f'Base        : {len(z):,} Zeus entities with a Definitive id ('
          + ' + '.join(f'{n:,} {k}' for k, n in counts) + ')')
    for k, n in e.Id_Status.value_counts().items():
        print(f'  {k:<52}: {n:,}')
    print(f'Tree        : {len(t):,} records = '
          f'{t.Node_Role.eq("Zeus id").sum():,} Zeus ids + '
          f'{t.Node_Role.eq("Owner only").sum():,} owners only '
          f'({len(ext):,} in no hierarchy view)')
    print(f'Trees       : {len(roots):,}  '
          f'({(roots.Descendant_Count > 0).sum():,} with more than one '
          f'member), max depth {t.Level.max()}')
    print(f'Entities    : {placed.Is_Root.eq(False).sum():,} of '
          f'{len(placed):,} placed have a Definitive owner above them; '
          f'{placed.Has_Intermediate_Parent.sum():,} through an intermediate')
    rows = checks(z, e, t, counts)
    print_checks(rows)
    wbp = workbook_path(prefix, 'Definitive_Only_Hierarchy')
    build_workbook(z, e, t, rows, wbp, os.path.basename(prefix), a.accuracy,
                   not a.no_imports)
    print(f'Wrote {prefix}_entities.csv  (+ _hierarchy.csv, _immediate.csv, '
          f'_ultimate.csv)')
    print(f'Wrote {wbp}')


if __name__ == '__main__':
    main()
